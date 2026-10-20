"""
生成初赛展示 PPT（16:9，深色，与产品界面同一套视觉）。

用法：python scripts/build_ppt.py
输出：deliverables/信任守护师_赛题2_初赛PPT.pptx
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "deliverables")
SHOT = os.path.join(os.path.dirname(ROOT), "")  # 截图在 Production 根目录

BG = RGBColor(0x0A, 0x0F, 0x1C)
PANEL = RGBColor(0x13, 0x1C, 0x30)
LINE = RGBColor(0x24, 0x36, 0x56)
WHITE = RGBColor(0xE8, 0xEE, 0xFB)
MUTED = RGBColor(0x8E, 0xA3, 0xC4)
DIM = RGBColor(0x5F, 0x73, 0x95)
GREEN = RGBColor(0x35, 0xD3, 0x9A)
AMBER = RGBColor(0xFF, 0xB0, 0x20)
RED = RGBColor(0xFF, 0x4D, 0x6D)
BLUE = RGBColor(0x4D, 0xA3, 0xFF)

FONT = "Microsoft YaHei"
W, H = Inches(13.333), Inches(7.5)


def shot(name: str) -> str:
    """截图统一放在 deliverables/screenshots/ 下。"""
    p = os.path.join(ROOT, "deliverables", "screenshots", name)
    if os.path.isfile(p):
        return p
    # 兼容旧路径
    for base in (ROOT,):
        q = os.path.join(base, name)
        if os.path.isfile(q):
            return q
    return ""


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def blank(prs: Presentation):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    return s


def text(slide, x, y, w, h, runs, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         line_spacing=1.35, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [[(runs, size, color, bold)]]
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        if isinstance(para, tuple):
            para, gap = para
            p.space_before = Pt(gap)
        for item in para:
            t, sz, col, bd = (list(item) + [size, color, bold])[:4]
            r = p.add_run()
            r.text = t
            r.font.size = Pt(sz)
            r.font.color.rgb = col
            r.font.bold = bd
            r.font.name = FONT
    return box


def kicker(slide, txt, color=BLUE):
    text(slide, 0.75, 0.55, 8, 0.35, txt, size=12, color=color, bold=True)


def title(slide, txt, sub=""):
    text(slide, 0.75, 0.92, 11.8, 0.9, txt, size=32, bold=True)
    if sub:
        text(slide, 0.75, 1.62, 11.5, 0.38, sub, size=13.5, color=MUTED)


def rule(slide, y=1.42):
    return  # 标题下方的绿色装饰线会与副标题打架，直接不画
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(y), Inches(1.2), Inches(0.045))
    ln.fill.solid()
    ln.fill.fore_color.rgb = GREEN
    ln.line.fill.background()


def panel(slide, x, y, w, h, color=PANEL, line=LINE, radius=0.06):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.color.rgb = line
    sh.line.width = Pt(0.75)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def pic(slide, path, x, y, w, h):
    if path and os.path.isfile(path):
        slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        panel(slide, x, y, w, h)
        text(slide, x, y + h / 2 - 0.2, w, 0.4, "（截图占位）", size=12, color=DIM, align=PP_ALIGN.CENTER)


def footer(slide, idx):
    text(slide, 0.75, 7.02, 6, 0.3, "信任守护师 · 欧莱雅第二届美妆科技黑客松 · 赛题2", size=9, color=DIM)
    text(slide, 11.6, 7.02, 1.0, 0.3, f"{idx:02d}", size=9, color=DIM, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- 各页

def s01_cover(prs):
    s = blank(prs)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.14), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()

    text(s, 0.9, 1.35, 8, 0.4, "欧莱雅第二届美妆科技黑客松 · 赛题 2", size=13, color=GREEN, bold=True)
    text(s, 0.9, 1.9, 9.6, 1.4, "信任守护师", size=58, bold=True)
    text(s, 0.9, 3.05, 10.4, 0.8, "让「眼见为实」重新变得可验证", size=26, color=RGBColor(0xC9, 0xD8, 0xEF))
    text(s, 0.9, 4.15, 9.8, 1.2,
         "一个基于像素级证据链的美妆内容核验 Agent：不只告诉你「像假的」，\n而是告诉你「哪一块、凭什么、违反了哪一条」。",
         size=14, color=MUTED, line_spacing=1.6)

    for i, (k, v, c) in enumerate([("五维", "取证维度", BLUE), ("3 场景", "统一架构", GREEN),
                                  ("< 4s", "本地 CPU 推理", AMBER), ("100%", "证据可回溯", RED)]):
        x = 0.9 + i * 2.25
        panel(s, x, 5.55, 2.0, 1.0)
        text(s, x, 5.68, 2.0, 0.4, k, size=19, color=c, bold=True, align=PP_ALIGN.CENTER)
        text(s, x, 6.12, 2.0, 0.3, v, size=10, color=DIM, align=PP_ALIGN.CENTER)

    text(s, 10.4, 6.55, 2.4, 0.3, "2026.10", size=11, color=DIM, align=PP_ALIGN.RIGHT)


def s02_pain(prs):
    s = blank(prs)
    kicker(s, "01  场景洞察")
    title(s, "一条笔记，可能是 AI 写、AI 修、AI 刷出来的", "美妆是内容信任危机最集中的行业：功效看不见、对比图最容易被修、评论区最容易被刷")
    rule(s)

    pic(s, shot("high.png"), 0.75, 2.18, 7.15, 4.25)

    items = [
        ("文案", "「7 天白 3 个度」——AI 一天能写 200 篇", RED),
        ("配图", "对比图局部磨皮、提亮、拼接，肉眼难辨", AMBER),
        ("评论", "200 条清一色好评，12 分钟内投放完毕", BLUE),
    ]
    y = 2.05
    for k, v, c in items:
        panel(s, 8.45, y, 4.15, 1.35)
        text(s, 8.62, y + 0.16, 0.6, 0.4, k, size=13, color=c, bold=True)
        text(s, 8.62, y + 0.55, 3.8, 0.7, v, size=12, color=RGBColor(0xC9, 0xD8, 0xEF), line_spacing=1.4)
        y += 1.52

    text(s, 8.45, 6.05, 4.15, 0.5, "平台审核员每天面对这些内容，可用工具却只有「举报」。",
         size=10.5, color=DIM, line_spacing=1.5)
    footer(s, 2)


def s03_insight(prs):
    s = blank(prs)
    kicker(s, "02  为什么现有方案不够")
    title(s, "平台要的不是一个概率，而是一条可以申诉的证据", "")
    rule(s)

    # 左：黑盒
    panel(s, 0.75, 1.95, 5.6, 3.95)
    text(s, 1.0, 2.2, 5.0, 0.4, "常规做法", size=12, color=DIM, bold=True)
    text(s, 1.0, 2.65, 5.0, 0.5, "「AIGC 概率 87%」", size=24, bold=True, color=RED)
    text(s, 1.0, 3.35, 5.0, 2.6,
         "· 审核员无法验证，也不敢据此下架\n"
         "· 创作者无法申诉，只能反复吵架\n"
         "· 跨生成器泛化差，模型一换就失灵\n"
         "· 检测器本身就在和生成器打军备竞赛",
         size=13, color=MUTED, line_spacing=1.9)

    # 右：证据链
    panel(s, 6.9, 1.95, 5.68, 3.95, color=RGBColor(0x0F, 0x22, 0x1C), line=RGBColor(0x27, 0x53, 0x44))
    text(s, 7.15, 2.2, 5.0, 0.4, "我们的做法", size=12, color=GREEN, bold=True)
    text(s, 7.15, 2.65, 5.2, 0.5, "「这一块，凭什么」", size=24, bold=True, color=GREEN)
    text(s, 7.15, 3.35, 5.2, 2.6,
         "· 每个结论挂载可核对的数值证据\n"
         "· 篡改区域可视化，边界可以放大看\n"
         "· 不确定时主动降级，而不是硬给高分\n"
         "· 直接输出可写入处置工单的违规条款",
         size=13, color=RGBColor(0xB8, 0xD8, 0xCC), line_spacing=1.9)

    text(s, 0.75, 6.55, 11.8, 0.4,
         "赛题原文：「可解释性完整度将作为方案可信度的重要评估依据」——这就是我们的技术主线。",
         size=11, color=DIM)
    footer(s, 3)


def s04_arch(prs):
    s = blank(prs)
    kicker(s, "03  技术方案")
    title(s, "五维证据链：像素 → 语义 → 行为 → 一致性 → 规则", "单点证据会误报，多层互证才敢下结论")
    rule(s)

    CYAN = RGBColor(0x3E, 0xC1, 0xD8)
    layers = [
        ("像素层", "ELA 重压缩误差 · 噪声残差 · 高频细节缺失（磨皮）\n频域幂律斜率 · SIFT 拷贝-移动 · EXIF", BLUE, 0.75),
        ("语义层", "9 项文体特征：句长变异、连接词密度、信息具体度、\n模板短语、用字丰富度…", RGBColor(0x7C, 0x6B, 0xF0), 2.36),
        ("行为层", "模板簇聚类 · 短时爆发比 · 情感一致性 ·\n账号画像（昵称熵/注册聚集）", GREEN, 3.97),
        ("一致性", "皮肤区 CIELAB 色差 vs 文案宣称\n夸大倍数 · 两图光照方向差", CYAN, 5.58),
        ("规则层", "《广告法》第九条、第二十八条\n《化妆品监督管理条例》\n第十六条、第四十三条", AMBER, 7.19),
    ]
    for name, desc, color, x in layers:
        panel(s, x, 2.18, 1.45, 1.95)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.18), Inches(1.45), Inches(0.07))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
        text(s, x + 0.10, 2.36, 1.28, 0.4, name, size=15, bold=True, color=color)
        text(s, x + 0.10, 2.80, 1.30, 1.25, desc, size=8.5, color=MUTED, line_spacing=1.45)

    # 融合
    panel(s, 8.85, 2.18, 3.7, 1.95, color=RGBColor(0x1A, 0x16, 0x0E), line=RGBColor(0x5A, 0x44, 0x1A))
    text(s, 9.05, 2.34, 3.3, 0.4, "证据融合", size=15, bold=True, color=AMBER)
    text(s, 9.05, 2.78, 3.3, 1.3,
         "加权融合 + 多信号互证\n确证性证据直接升级\n覆盖不足自动降级",
         size=10, color=RGBColor(0xD8, 0xC8, 0xA8), line_spacing=1.6)

    # 结论条
    panel(s, 0.75, 4.42, 11.8, 0.95, color=RGBColor(0x0F, 0x1A, 0x2B))
    text(s, 1.0, 4.62, 11.3, 0.7,
         [[("输出：", 13, MUTED, True), ("风险分级（高/中/低）", 13, WHITE, True),
           ("  ·  ", 13, DIM, False), ("判据清单", 13, WHITE, True),
           ("  ·  ", 13, DIM, False), ("篡改区域热力图", 13, WHITE, True),
           ("  ·  ", 13, DIM, False), ("违规条款引用", 13, WHITE, True),
           ("  ·  ", 13, DIM, False), ("处置建议", 13, WHITE, True)]],
         line_spacing=1.5)

    text(s, 0.75, 5.58, 11.8, 1.3,
         "固定 DAG 编排（非自由 ReAct）：场景识别 → 并行取证 → 融合定级 → 生成结论，"
         "演示现场不会出现循环调用、幻觉工具、耗时不可控。\n"
         "为什么不上大模型做判定：黑盒模型给不出可核对的数值证据，跨生成器泛化也差；"
         "本项目坚持「每个结论都能反推」，大模型（Qwen2.5-VL）放在第二期做语义复核。",
         size=10.5, color=MUTED, line_spacing=1.6)
    footer(s, 4)


def s05_pixel(prs):
    s = blank(prs)
    kicker(s, "04  能力演示 · 像素层")
    title(s, "篡改区域定位：热图落在磨皮的地方", "系统自己找到区域、自己画出边界，并给出每个维度的数值")
    rule(s)

    pic(s, shot("high.png"), 0.75, 2.16, 8.05, 4.78)

    panel(s, 9.15, 2.16, 3.4, 4.78)
    rows = [("噪声残差", "21.4% 区域异常", AMBER),
            ("高频细节", "-4.6σ 低于周边", RED),
            ("重压缩误差", "z = 9.33", RED),
            ("图文色差", "ΔE 3.81 vs 宣称 10.5", RED)]
    y = 2.25
    for k, v, c in rows:
        text(s, 9.4, y, 2.9, 0.3, k, size=10, color=DIM)
        text(s, 9.4, y + 0.3, 2.9, 0.4, v, size=13, bold=True, color=c)
        y += 1.1
    text(s, 9.4, 6.25, 2.9, 0.5, "单次分析 2–4 秒\n纯 CPU", size=9, color=DIM, line_spacing=1.5)
    footer(s, 5)


def s06_consistency(prs):
    s = blank(prs)
    kicker(s, "05  能力演示 · 图文一致性")
    title(s, "把「美白三个度」换算成可以复算的数字", "肤色转到 CIELAB 空间测色差，与文案宣称直接对照")
    rule(s)

    # 公式卡
    panel(s, 0.75, 2.14, 3.85, 1.2)
    text(s, 0.95, 2.26, 3.5, 0.35, "实测色差", size=10, color=DIM)
    text(s, 0.95, 2.60, 3.5, 0.55, "ΔE = 3.81", size=26, bold=True, color=GREEN)

    panel(s, 4.8, 2.14, 3.85, 1.2)
    text(s, 5.0, 2.26, 3.5, 0.35, "文案宣称折算", size=10, color=DIM)
    text(s, 5.0, 2.60, 3.5, 0.55, "ΔE ≈ 10.5", size=26, bold=True, color=RED)

    panel(s, 8.85, 2.14, 3.7, 1.2, color=RGBColor(0x2A, 0x12, 0x18), line=RGBColor(0x6B, 0x22, 0x33))
    text(s, 9.05, 2.26, 3.3, 0.35, "夸大倍数", size=10, color=DIM)
    text(s, 9.05, 2.60, 3.3, 0.55, "2.75 ×", size=26, bold=True, color=RED)

    text(s, 0.75, 3.62, 11.8, 0.9,
         "实测：两图皮肤区域平均 L*a*b* 差的欧氏距离。宣称：从文案抽取「白 N 个度」，按一个色号 ≈ ΔE 3.5 折算。\n"
         "这不是「我觉得像假的」，而是一个可以拿给创作者看的数字。",
         size=12, color=MUTED, line_spacing=1.75)

    # 对照组
    panel(s, 0.75, 4.78, 5.75, 1.75, color=RGBColor(0x0F, 0x22, 0x1C), line=RGBColor(0x27, 0x53, 0x44))
    text(s, 1.0, 4.92, 5.3, 0.35, "对照组：真实测评", size=11, color=GREEN, bold=True)
    text(s, 1.0, 5.34, 5.3, 1.0, "同一场景、无局部修改、文案无夸大\n→ 图文一致性 0.00，整体 低风险 0.14",
         size=12, color=RGBColor(0xB8, 0xD8, 0xCC), line_spacing=1.65)

    panel(s, 6.9, 4.78, 5.65, 1.75, color=RGBColor(0x2A, 0x12, 0x18), line=RGBColor(0x6B, 0x22, 0x33))
    text(s, 7.15, 4.92, 5.2, 0.35, "实验组：高仿对比图", size=11, color=RED, bold=True)
    text(s, 7.15, 5.34, 5.2, 1.0, "局部磨皮 + 提亮 + 拼接，文案宣称 3 个度\n→ 整体 高风险 0.87",
         size=12, color=RGBColor(0xE0, 0xB8, 0xC0), line_spacing=1.65)
    footer(s, 6)


def s07_comments(prs):
    s = blank(prs)
    kicker(s, "06  能力演示 · 评论区")
    title(s, "不看单条评论，看「这批账号在做什么」", "12 条同一模板的评论，8 分钟内投放完毕")
    rule(s)

    pic(s, shot("comments.png"), 0.75, 2.16, 8.05, 4.78)

    panel(s, 9.35, 2.0, 3.2, 4.93)
    rows = [("模板簇", "12 条", RED), ("平均相似度", "0.92", RED),
            ("短时爆发比", "67%", AMBER), ("异常账号", "user8821 等", AMBER)]
    y = 2.25
    for k, v, c in rows:
        text(s, 9.6, y, 2.7, 0.3, k, size=10, color=DIM)
        text(s, 9.6, y + 0.3, 2.7, 0.4, v, size=14, bold=True, color=c)
        y += 1.1
    text(s, 9.4, 6.2, 2.9, 0.5, "确证性证据 → 直接高风险", size=9, color=DIM, line_spacing=1.5)
    footer(s, 7)


def s08_rules(prs):
    s = blank(prs)
    kicker(s, "07  能力演示 · 合规规则")
    title(s, "输出的不是概率，是可以写进工单的违规条目", "")
    rule(s)

    rows = [
        ("绝对化用语", "「全网最低价」", "《广告法》第九条", "不得使用国家级、最高级、最佳等用语"),
        ("特殊化妆品未标批号", "「美白」", "《化妆品监督管理条例》第十六条", "美白祛斑等须标注国妆特字注册编号"),
        ("量化功效承诺缺乏依据", "「立竿见影」", "《化妆品功效宣称评价规范》", "功效宣称应有评价依据"),
        ("宣称医疗作用", "「消炎」「药妆」", "《化妆品监督管理条例》第四十三条", "化妆品不得明示或暗示医疗作用"),
    ]
    y = 2.05
    for name, hit, clause, note in rows:
        panel(s, 0.75, y, 11.8, 0.95)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(y), Inches(0.06), Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RED
        bar.line.fill.background()
        text(s, 1.0, y + 0.14, 2.6, 0.35, name, size=12.5, bold=True)
        text(s, 1.0, y + 0.55, 2.6, 0.35, hit, size=11, color=AMBER)
        text(s, 3.75, y + 0.14, 3.9, 0.35, clause, size=10.5, color=BLUE)
        text(s, 3.75, y + 0.55, 8.5, 0.35, note, size=10, color=MUTED)
        y += 1.05

    text(s, 0.75, 6.45, 11.8, 0.4,
         "规则层零算力、零幻觉，却是「落地价值」最直接的论据：它输出的是处置依据，不是置信度。",
         size=11, color=DIM)
    footer(s, 8)


def s09_product(prs):
    s = blank(prs)
    kicker(s, "08  产品形态")
    title(s, "三栏式工作台 + 对话追问", "左边输入、中间证据、右边结论；所有证据可点击联动")
    rule(s)

    pic(s, shot("low.png"), 0.75, 2.16, 8.05, 4.78)

    panel(s, 9.15, 2.16, 3.4, 2.25)
    text(s, 9.4, 2.30, 3.0, 0.35, "交互设计", size=12, bold=True, color=GREEN)
    text(s, 9.4, 2.68, 3.0, 1.6,
         "· 正片/热图透明度滑杆\n· 4 个维度热图切换\n· 证据卡 ↔ 区域框联动\n· 对话引用证据回答",
         size=10, color=MUTED, line_spacing=1.75)

    panel(s, 9.15, 4.55, 3.4, 2.39, color=RGBColor(0x0F, 0x1A, 0x2B))
    text(s, 9.4, 4.70, 3.0, 0.35, "工程约束", size=12, bold=True, color=BLUE)
    text(s, 9.4, 5.08, 3.0, 1.75,
         "· 纯 CPU、可离线\n· 无前端构建，双击即用\n· 单次 2–4 秒\n· 模型可替换、可降级",
         size=10, color=MUTED, line_spacing=1.75)
    footer(s, 9)


def s10_dataset(prs):
    s = blank(prs)
    kicker(s, "09  测试数据包", GREEN)
    title(s, "我们连测试数据一起交：50 条，覆盖三场景与三档难度",
          "赛题看的不只是跑分，而是这份数据本身经不经得起看")
    rule(s)

    # 左：构成
    panel(s, 0.75, 2.18, 5.9, 3.15)
    text(s, 1.0, 2.30, 5.4, 0.35, "样本构成", size=11.5, color=DIM, bold=True)
    rows = [
        ("场景", "种草 20　·　评论 15　·　素材 15"),
        ("真伪", "伪造 28　·　真实 22"),
        ("难度", "easy 7　·　medium 17　·　hard 4　·　real 22"),
        ("标注", "12 张 mask.png（逐像素篡改区域）"),
    ]
    y = 2.75
    for k, v in rows:
        text(s, 1.0, y, 1.0, 0.3, k, size=10, color=MUTED)
        text(s, 1.75, y, 4.7, 0.3, v, size=10.5, color=WHITE, bold=True)
        y += 0.62

    # 右：伪造手法
    panel(s, 6.85, 2.18, 5.7, 3.15, color=RGBColor(0x0F, 0x1A, 0x2B))
    text(s, 7.1, 2.30, 5.2, 0.35, "五类伪造手法（不是同一种东西做五遍）", size=11.5, color=DIM, bold=True)
    kinds = [
        ("大面积磨皮 + 提亮", "4 条"),
        ("局部磨皮 + 拼接", "4 条"),
        ("轻度磨皮 + 拼接 + 重压缩", "4 条"),
        ("重度平滑 + 色彩拉伸（模拟 AIGC 重绘）", "8 条"),
        ("模板化评论 + 短时间集中投放", "8 条"),
    ]
    y = 2.75
    for name, cnt in kinds:
        text(s, 7.1, y, 4.3, 0.3, "· " + name, size=10, color=RGBColor(0xC9, 0xD8, 0xEF))
        text(s, 11.35, y, 0.9, 0.3, cnt, size=10, color=AMBER, bold=True, align=PP_ALIGN.RIGHT)
        y += 0.48

    # 底：目录结构 + 说明
    panel(s, 0.75, 5.50, 7.4, 1.15, color=RGBColor(0x0B, 0x14, 0x22))
    text(s, 0.95, 5.60, 7.0, 0.95,
         "sample_001/\n"
         "  ├ meta.txt     场景 / 真伪 / 难度 / 伪造方式\n"
         "  ├ mask.png     篡改区域（仅图片类样本）\n"
         "  └ post/        content.txt + image_1.jpg + image_2.jpg",
         size=8.5, color=MUTED, line_spacing=1.5)

    text(s, 8.35, 5.62, 4.2, 1.0,
         "每条样本都能复现：\n把 meta.txt 的字段喂给 eval/run_eval.py，\n就能重算出 PPT 里的所有指标。",
         size=9.5, color=RGBColor(0x9F, 0xC8, 0xFF), line_spacing=1.6)
    footer(s, 10)


def s10_eval(prs):
    s = blank(prs)
    kicker(s, "10  效果与严谨性")
    title(s, "我们交的不只是系统，还有一套可复现的评测", "50 条自建样本，真伪 28:22，带 ground-truth mask")
    rule(s)

    stats = [("50", "测试样本", BLUE), ("28 / 22", "伪造 / 真实", AMBER),
             ("0.77", "Precision", GREEN), ("0.46", "Recall", RED), ("0.23", "定位 IoU", RGBColor(0x7C, 0x6B, 0xF0))]
    x = 0.75
    for k, v, c in stats:
        panel(s, x, 2.12, 2.22, 1.15)
        text(s, x, 2.22, 2.22, 0.5, k, size=20, bold=True, color=c, align=PP_ALIGN.CENTER)
        text(s, x, 2.78, 2.22, 0.3, v, size=9.5, color=DIM, align=PP_ALIGN.CENTER)
        x += 2.4

    # 分场景
    panel(s, 0.75, 3.52, 5.9, 1.6)
    text(s, 1.0, 3.6, 5.4, 0.35, "分场景 F1", size=11, color=DIM, bold=True)
    text(s, 1.0, 4.0, 5.4, 1.0,
         "种草图文 0.63　·　评论区 0.93　·　AI 素材 0.00",
         size=15, bold=True, color=WHITE)
    text(s, 1.0, 4.62, 5.4, 0.4,
         "素材场景在自建集上没有区分度，只作辅助筛查，见下一页",
         size=9.5, color=RGBColor(0x8A, 0x9B, 0xB8))

    panel(s, 6.85, 3.52, 5.7, 1.6, color=RGBColor(0x2A, 0x12, 0x18), line=RGBColor(0x6B, 0x22, 0x33))
    text(s, 7.1, 3.6, 5.2, 0.35, "已知短板（不掩饰）", size=11, color=RED, bold=True)
    text(s, 7.1, 4.0, 5.2, 1.0,
         "早期 F1 是 0.69，其中素材场景的检出\n来自 ELA 的 z 分数退化（z=876445）；\n修正后如实回落到 0.58。",
         size=10, color=RGBColor(0xE0, 0xB8, 0xC0), line_spacing=1.5)

    text(s, 0.75, 5.35, 11.8, 1.35,
         "评测方法：样本级 Precision / Recall 以 0.5 分为阈值；定位 IoU 用系统输出的可疑区域二值化后与 mask.png 对比。\n"
         "数据包里同时包含真实样本，就是为了检验系统会不会「见谁都像假的」——正是这一步暴露了上面的误报。",
         size=11, color=MUTED, line_spacing=1.75)
    footer(s, 11)


def s11_value(prs):
    s = blank(prs)
    kicker(s, "11  落地价值")
    title(s, "从「事后举报」变成「事前拦截」", "")
    rule(s)

    steps = [("抽检", "每天 5% 内容\n自动过一遍"), ("核验", "五维证据链\n2–4 秒出结论"),
             ("分级", "高危下架\n中危限流"), ("申诉", "证据包直接\n回传创作者")]
    x = 0.75
    for i, (k, v) in enumerate(steps):
        panel(s, x, 2.12, 2.55, 1.55)
        text(s, x + 0.2, 2.26, 2.2, 0.3, f"0{i+1}", size=11, color=GREEN, bold=True)
        text(s, x + 0.2, 2.56, 2.2, 0.35, k, size=17, bold=True)
        text(s, x + 0.2, 3.0, 2.2, 0.6, v, size=10, color=MUTED, line_spacing=1.45)
        x += 2.78

    cols = [("对平台", "审核效率提升，处置有据可依，减少误伤与舆情", BLUE),
            ("对品牌", "保护真实口碑不被水军稀释，掌握竞品造假证据", GREEN),
            ("对消费者", "看到的内容有可信度标注，知道该不该信", AMBER)]
    x = 0.75
    for k, v, c in cols:
        panel(s, x, 4.10, 3.85, 1.5)
        text(s, x + 0.22, 4.24, 3.4, 0.35, k, size=14, bold=True, color=c)
        text(s, x + 0.22, 4.66, 3.4, 0.85, v, size=11, color=MUTED, line_spacing=1.6)
        x += 4.02

    text(s, 0.75, 5.95, 11.8, 0.8,
         "合规是刚需：市场监管总局对化妆品功效宣称的专项整治一直在进行，平台侧的审核压力是真实存在的。",
         size=11, color=DIM, line_spacing=1.6)
    footer(s, 12)


def s12_roadmap(prs):
    s = blank(prs)
    kicker(s, "12  未来规划")
    title(s, "三条线把这个 Demo 变成产品", "")
    rule(s)

    items = [
        ("1–3 个月", "模型增强", "接入开源预训练 AIGC 检测器做集成；\n引入 Qwen2.5-VL 做语义层复核；\n用真实平台数据做阈值标定。", BLUE),
        ("3–6 个月", "场景扩展", "从美妆扩到母婴、保健品等功效宣称密集品类；\n从图文扩到短视频帧级核验。", GREEN),
        ("6–12 个月", "平台化", "对接内容中台，做成审核工作流的插件；\n证据包标准化，支持监管报送。", AMBER),
    ]
    x = 0.75
    for k, name, desc, c in items:
        panel(s, x, 2.12, 3.85, 3.25)
        text(s, x + 0.25, 2.30, 3.3, 0.3, k, size=11, color=c, bold=True)
        text(s, x + 0.25, 2.68, 3.3, 0.45, name, size=19, bold=True)
        text(s, x + 0.25, 3.28, 3.35, 1.95, desc, size=11, color=MUTED, line_spacing=1.85)
        x += 4.02

    panel(s, 0.75, 5.62, 11.8, 1.0, color=RGBColor(0x0F, 0x22, 0x1C), line=RGBColor(0x27, 0x53, 0x44))
    text(s, 1.0, 5.78, 11.3, 0.7,
         [[("我们相信：", 13, GREEN, True),
           ("「内容可信度」会成为和「内容安全」同等重要的基础设施——而它的前提，是每一个判断都能被解释。", 13, WHITE, False)]],
         line_spacing=1.5)
    footer(s, 13)


def main() -> None:
    os.makedirs(OUT_DIR, exist_ok=True)
    prs = new_deck()
    s01_cover(prs)
    s02_pain(prs)
    s03_insight(prs)
    s04_arch(prs)
    s05_pixel(prs)
    s06_consistency(prs)
    s07_comments(prs)
    s08_rules(prs)
    s09_product(prs)
    s10_dataset(prs)
    s10_eval(prs)
    s11_value(prs)
    s12_roadmap(prs)

    out = os.path.join(OUT_DIR, "信任守护师_赛题2_初赛PPT.pptx")
    prs.save(out)
    print("saved:", out, os.path.getsize(out) // 1024, "KB")


if __name__ == "__main__":
    main()
